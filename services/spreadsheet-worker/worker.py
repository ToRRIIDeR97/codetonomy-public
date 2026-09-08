#!/usr/bin/env python3
"""Isolated, schema-driven spreadsheet worker for Codetonomy."""

from __future__ import annotations

import json
import math
import os
import sys
import tempfile
from pathlib import Path
from typing import Any
from zipfile import ZipFile

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill

MAX_INPUT_BYTES = 512 * 1024
MAX_WORKBOOK_BYTES = 20 * 1024 * 1024
MAX_EXPANDED_BYTES = 32 * 1024 * 1024
MAX_WORKBOOK_CELLS = 250_000
REQUIRED_SHEETS = {"Summary", "Historical", "Assumptions", "Valuation", "Sources"}


def fail(message: str) -> None:
    print(json.dumps({"ok": False, "error": message}))


def finite(value: Any, label: str) -> float:
    if isinstance(value, bool) or not isinstance(value, (int, float)) or not math.isfinite(value):
        raise ValueError(f"{label} must be a finite number")
    return float(value)


def safe_text(value: Any, label: str, maximum: int = 500) -> str:
    if not isinstance(value, str) or not value.strip() or len(value) > maximum or any(ord(char) < 32 and char not in "\t" for char in value):
        raise ValueError(f"{label} must be non-empty text of at most {maximum} characters")
    text = value.strip()
    return f"'{text}" if text[0] in "=+-@" else text


def confined(workspace: str, requested: str, must_exist: bool) -> tuple[Path, str]:
    root = Path(workspace).resolve(strict=True)
    raw = Path(requested)
    if raw.is_absolute() or ".." in raw.parts or not raw.parts:
        raise ValueError("Workbook path is outside the workspace")
    target = root.joinpath(raw)
    ancestor = target if must_exist else target.parent
    while not ancestor.exists():
        if must_exist or ancestor == root.parent:
            raise ValueError("Workbook path does not exist")
        ancestor = ancestor.parent
    if not ancestor.resolve(strict=True).is_relative_to(root):
        raise ValueError("Workbook path resolves outside the workspace")
    if must_exist:
        resolved = target.resolve(strict=True)
        if not resolved.is_relative_to(root) or not resolved.is_file() or target.is_symlink():
            raise ValueError("Workbook must be a regular workspace file")
        if resolved.stat().st_size > MAX_WORKBOOK_BYTES:
            raise ValueError("Workbook exceeds 20 MiB")
    return target, raw.as_posix()


def style_header(sheet: Any, row: int = 1) -> None:
    fill = PatternFill("solid", fgColor="DCEAF7")
    for cell in sheet[row]:
        cell.font = Font(bold=True, color="16324F")
        cell.fill = fill
        cell.alignment = Alignment(horizontal="center")


def validate_spec(spec: Any) -> dict[str, Any]:
    if not isinstance(spec, dict):
        raise ValueError("spec must be an object")
    company = safe_text(spec.get("company"), "company", 160)
    output_path = safe_text(spec.get("outputPath"), "outputPath", 1024)
    historical = spec.get("historical")
    scenarios = spec.get("scenarios")
    sources = spec.get("sources")
    if not isinstance(historical, list) or not 1 <= len(historical) <= 50:
        raise ValueError("historical must contain 1-50 rows")
    if not isinstance(scenarios, list) or len(scenarios) != 3:
        raise ValueError("scenarios must contain exactly three scenarios")
    if not isinstance(sources, list) or not 1 <= len(sources) <= 100:
        raise ValueError("sources must contain 1-100 notes")
    clean_historical = []
    years: set[int] = set()
    for index, row in enumerate(historical):
        if not isinstance(row, dict):
            raise ValueError(f"historical[{index}] must be an object")
        year_value = finite(row.get("year"), f"historical[{index}].year")
        year = int(year_value)
        if year != year_value or year < 1900 or year > 2200 or year in years:
            raise ValueError("historical years must be unique integers from 1900-2200")
        years.add(year)
        clean_historical.append({
            "year": year,
            "revenue": finite(row.get("revenue"), f"historical[{index}].revenue"),
            "ebitda": finite(row.get("ebitda"), f"historical[{index}].ebitda"),
            "freeCashFlow": finite(row.get("freeCashFlow"), f"historical[{index}].freeCashFlow"),
            "source": safe_text(row["source"], f"historical[{index}].source") if row.get("source") else "See Sources",
        })
    clean_historical.sort(key=lambda item: item["year"])
    clean_scenarios = []
    names: set[str] = set()
    for index, scenario in enumerate(scenarios):
        if not isinstance(scenario, dict):
            raise ValueError(f"scenarios[{index}] must be an object")
        name = safe_text(scenario.get("name"), f"scenarios[{index}].name", 40)
        if name in names:
            raise ValueError("scenario names must be unique")
        names.add(name)
        growth = finite(scenario.get("revenueGrowth"), f"scenarios[{index}].revenueGrowth")
        margin = finite(scenario.get("ebitdaMargin"), f"scenarios[{index}].ebitdaMargin")
        if not -0.5 <= growth <= 1 or not -1 <= margin <= 1:
            raise ValueError("scenario rates are outside safe bounds")
        clean_scenarios.append({"name": name, "revenueGrowth": growth, "ebitdaMargin": margin})
    clean_sources = []
    for index, source in enumerate(sources):
        if not isinstance(source, dict):
            raise ValueError(f"sources[{index}] must be an object")
        page = source.get("page")
        if page is not None and (isinstance(page, bool) or not isinstance(page, int) or page < 1):
            raise ValueError(f"sources[{index}].page must be a positive integer")
        clean_sources.append({
            "label": safe_text(source.get("label"), f"sources[{index}].label"),
            "source": safe_text(source.get("source"), f"sources[{index}].source", 2000),
            "page": page,
        })
    discount = finite(spec.get("discountRate"), "discountRate")
    terminal = finite(spec.get("terminalGrowthRate"), "terminalGrowthRate")
    tax = finite(spec.get("taxRate", 0.21), "taxRate")
    if not 0 < discount <= 1 or not -0.2 <= terminal < discount or not 0 <= tax <= 1:
        raise ValueError("discount, terminal growth, or tax rate is invalid")
    return {
        "company": company,
        "outputPath": output_path,
        "historical": clean_historical,
        "scenarios": clean_scenarios,
        "discountRate": discount,
        "terminalGrowthRate": terminal,
        "taxRate": tax,
        "sources": clean_sources,
    }


def create_valuation(workspace: str, raw_spec: Any) -> dict[str, Any]:
    spec = validate_spec(raw_spec)
    target, relative_path = confined(workspace, spec["outputPath"], False)
    if target.suffix.lower() != ".xlsx":
        raise ValueError("outputPath must end in .xlsx")
    target.parent.mkdir(parents=True, exist_ok=True)
    workbook = Workbook()
    summary = workbook.active
    summary.title = "Summary"
    historical = workbook.create_sheet("Historical")
    assumptions = workbook.create_sheet("Assumptions")
    valuation = workbook.create_sheet("Valuation")
    sources = workbook.create_sheet("Sources")

    summary.append(["Company", spec["company"]])
    summary.append(["Model", "Five-year discounted cash flow"])
    summary.append(["Status", "Formula-driven; verify before investment use"])
    summary["A1"].font = Font(bold=True)
    summary.column_dimensions["A"].width = 22
    summary.column_dimensions["B"].width = 48

    historical.append(["Year", "Revenue", "EBITDA", "Free Cash Flow", "Source"])
    for row in spec["historical"]:
        historical.append([row["year"], row["revenue"], row["ebitda"], row["freeCashFlow"], row["source"]])
    style_header(historical)
    historical.freeze_panes = "A2"
    historical.column_dimensions["E"].width = 50
    for row in historical.iter_rows(min_row=2, min_col=2, max_col=4):
        for cell in row:
            cell.number_format = "#,##0.00"

    assumptions.append(["Metric", *[scenario["name"] for scenario in spec["scenarios"]]])
    assumptions.append(["Revenue growth", *[scenario["revenueGrowth"] for scenario in spec["scenarios"]]])
    assumptions.append(["EBITDA margin", *[scenario["ebitdaMargin"] for scenario in spec["scenarios"]]])
    assumptions.append(["Discount rate", *[spec["discountRate"]] * 3])
    assumptions.append(["Terminal growth", *[spec["terminalGrowthRate"]] * 3])
    assumptions.append(["Tax rate", *[spec["taxRate"]] * 3])
    style_header(assumptions)
    for row in assumptions.iter_rows(min_row=2, min_col=2, max_col=4):
        for cell in row:
            cell.number_format = "0.0%"

    last_row = historical.max_row
    last_year = spec["historical"][-1]["year"]
    valuation.append(["Scenario", "Metric", *[last_year + year for year in range(1, 6)], "Terminal Value", "Enterprise Value"])
    style_header(valuation)
    row = 2
    for scenario_index, scenario in enumerate(spec["scenarios"], start=2):
        start = row
        valuation.cell(row, 1, scenario["name"])
        valuation.cell(row, 2, "Revenue")
        valuation.cell(row, 3, f"=Historical!B{last_row}*(1+Assumptions!{chr(64 + scenario_index)}2)")
        for column in range(4, 8):
            valuation.cell(row, column, f"={valuation.cell(row, column - 1).coordinate}*(1+Assumptions!{chr(64 + scenario_index)}2)")
        row += 1
        valuation.cell(row, 2, "EBITDA")
        for column in range(3, 8):
            valuation.cell(row, column, f"={valuation.cell(start, column).coordinate}*Assumptions!{chr(64 + scenario_index)}3")
        row += 1
        valuation.cell(row, 2, "Free Cash Flow")
        for column in range(3, 8):
            valuation.cell(row, column, f"={valuation.cell(start + 1, column).coordinate}*(1-Assumptions!{chr(64 + scenario_index)}6)")
        row += 1
        valuation.cell(row, 2, "Present Value")
        for year_index, column in enumerate(range(3, 8), start=1):
            valuation.cell(row, column, f"={valuation.cell(start + 2, column).coordinate}/(1+Assumptions!{chr(64 + scenario_index)}4)^{year_index}")
        valuation.cell(row, 8, f"={valuation.cell(start + 2, 7).coordinate}*(1+Assumptions!{chr(64 + scenario_index)}5)/(Assumptions!{chr(64 + scenario_index)}4-Assumptions!{chr(64 + scenario_index)}5)")
        valuation.cell(row, 9, f"=SUM(C{row}:G{row})+H{row}/(1+Assumptions!{chr(64 + scenario_index)}4)^5")
        row += 2
    valuation.freeze_panes = "C2"
    valuation.column_dimensions["B"].width = 20
    for row_cells in valuation.iter_rows(min_row=2, min_col=3, max_col=9):
        for cell in row_cells:
            cell.number_format = "#,##0.00"

    sources.append(["Label", "Source", "Page"])
    for source in spec["sources"]:
        sources.append([source["label"], source["source"], source["page"]])
    style_header(sources)
    sources.freeze_panes = "A2"
    sources.column_dimensions["A"].width = 28
    sources.column_dimensions["B"].width = 90

    workbook.calculation.fullCalcOnLoad = True
    workbook.calculation.forceFullCalc = True
    fd, temporary_name = tempfile.mkstemp(prefix=f".{target.name}.", suffix=".tmp", dir=target.parent)
    os.close(fd)
    try:
        workbook.save(temporary_name)
        os.replace(temporary_name, target)
        os.chmod(target, 0o644)
    finally:
        if os.path.exists(temporary_name):
            os.unlink(temporary_name)
    return inspect_workbook(workspace, relative_path)


def scan_workbook(workspace: str, path: str) -> tuple[dict[str, Any], dict[str, dict[str, Any]]]:
    target, relative_path = confined(workspace, path, True)
    with ZipFile(target) as archive:
        entries = archive.infolist()
        if len(entries) > 10_000 or sum(entry.file_size for entry in entries) > MAX_EXPANDED_BYTES:
            raise ValueError("Expanded workbook exceeds the inspection budget")
    workbook = load_workbook(target, read_only=True, data_only=False, keep_links=False)
    try:
        total = 0
        for sheet in workbook.worksheets:
            if sheet.max_row is None or sheet.max_column is None:
                raise ValueError("Worksheet dimensions are required for bounded inspection")
            total += sheet.max_row * sheet.max_column
            if total > MAX_WORKBOOK_CELLS:
                raise ValueError(f"Workbook exceeds {MAX_WORKBOOK_CELLS} inspected cells")
        values = {}
        for sheet in workbook.worksheets:
            values[sheet.title] = {cell.coordinate: cell.value for row in sheet.iter_rows()
                                   for cell in row if cell.value is not None}
        formulas = sum(isinstance(value, str) and value.startswith("=")
                       for sheet in values.values() for value in sheet.values())
        assumptions = values.get("Assumptions", {})
        inspection = {
            "path": relative_path,
            "sheets": [{"name": sheet.title, "rows": sheet.max_row, "columns": sheet.max_column} for sheet in workbook.worksheets],
            "formulas": formulas,
            "scenarios": [str(assumptions.get(f"{column}1", "")) for column in "BCD"] if assumptions else [],
        }
        return inspection, values
    finally:
        workbook.close()


def inspect_workbook(workspace: str, path: str) -> dict[str, Any]:
    return scan_workbook(workspace, path)[0]


def check_dcf(inspection: dict[str, Any], values: dict[str, dict[str, Any]]) -> dict[str, float]:
    # ponytail: verify the generated five-year DCF schema; other models need their own verifier.
    historical = values["Historical"]
    assumptions = values["Assumptions"]
    valuation = values["Valuation"]
    rows = next(sheet["rows"] for sheet in inspection["sheets"] if sheet["name"] == "Historical")
    source_rows = next(sheet["rows"] for sheet in inspection["sheets"] if sheet["name"] == "Sources")
    spec = validate_spec({
        "company": values["Summary"].get("B1"), "outputPath": inspection["path"],
        "historical": [{"year": historical.get(f"A{row}"), "revenue": historical.get(f"B{row}"),
                        "ebitda": historical.get(f"C{row}"), "freeCashFlow": historical.get(f"D{row}")}
                       for row in range(2, rows + 1)],
        "scenarios": [{"name": assumptions.get(f"{column}1"), "revenueGrowth": assumptions.get(f"{column}2"),
                       "ebitdaMargin": assumptions.get(f"{column}3")} for column in "BCD"],
        "discountRate": assumptions.get("B4"), "terminalGrowthRate": assumptions.get("B5"), "taxRate": assumptions.get("B6"),
        "sources": [{"label": values["Sources"].get(f"A{row}"), "source": values["Sources"].get(f"B{row}")}
                    for row in range(2, source_rows + 1)],
    })
    if values["Summary"].get("B2") != "Five-year discounted cash flow":
        raise ValueError("Unsupported valuation model")
    if [historical.get(f"A{row}") for row in range(2, rows + 1)] != [row["year"] for row in spec["historical"]]:
        raise ValueError("Historical years must be chronological")
    expected = {}
    results = {}
    for index, column in enumerate("BCD"):
        start = 2 + index * 5
        scenario = spec["scenarios"][index]
        if valuation.get(f"A{start}") != scenario["name"]:
            raise ValueError("Valuation scenario does not match assumptions")
        for row in (4, 5, 6):
            if assumptions.get(f"{column}{row}") != assumptions.get(f"B{row}"):
                raise ValueError("Discount, terminal growth and tax rates must match across scenarios")
        for year, forecast in enumerate("CDEFG", start=1):
            previous = f"Historical!B{rows}" if year == 1 else f"{chr(ord(forecast) - 1)}{start}"
            expected[f"{forecast}{start}"] = f"={previous}*(1+Assumptions!{column}2)"
            expected[f"{forecast}{start + 1}"] = f"={forecast}{start}*Assumptions!{column}3"
            expected[f"{forecast}{start + 2}"] = f"={forecast}{start + 1}*(1-Assumptions!{column}6)"
            expected[f"{forecast}{start + 3}"] = f"={forecast}{start + 2}/(1+Assumptions!{column}4)^{year}"
        expected[f"H{start + 3}"] = f"=G{start + 2}*(1+Assumptions!{column}5)/(Assumptions!{column}4-Assumptions!{column}5)"
        expected[f"I{start + 3}"] = f"=SUM(C{start + 3}:G{start + 3})+H{start + 3}/(1+Assumptions!{column}4)^5"
        revenue = spec["historical"][-1]["revenue"]
        cash_flows = [finite(revenue * (1 + scenario["revenueGrowth"]) ** year
                            * scenario["ebitdaMargin"] * (1 - spec["taxRate"]), "Forecast cash flow") for year in range(1, 6)]
        present = sum(cash / (1 + spec["discountRate"]) ** year for year, cash in enumerate(cash_flows, start=1))
        terminal = cash_flows[-1] * (1 + spec["terminalGrowthRate"]) / (spec["discountRate"] - spec["terminalGrowthRate"])
        results[scenario["name"]] = finite(present + terminal / (1 + spec["discountRate"]) ** 5, "Enterprise value")
    actual = {(sheet, address): value for sheet, cells in values.items() for address, value in cells.items()
              if isinstance(value, str) and value.startswith("=")}
    if actual != {("Valuation", address): formula for address, formula in expected.items()}:
        raise ValueError("Required DCF formulas are missing, altered, or contain unsupported references")
    return results


def verify_valuation(workspace: str, path: str) -> dict[str, Any]:
    inspection, values = scan_workbook(workspace, path)
    sheet_names = set(values)
    scenarios = inspection["scenarios"]
    calculation_error = None
    calculated = {}
    try:
        calculated = check_dcf(inspection, values)
    except (KeyError, ValueError, OverflowError, ZeroDivisionError) as error:
        calculation_error = str(error)
    checks = [
        {"id": "workbook-open", "passed": True, "message": "Workbook opens within inspection limits"},
        {"id": "required-sheet-check", "passed": REQUIRED_SHEETS.issubset(sheet_names), "message": "Required sheets are present" if REQUIRED_SHEETS.issubset(sheet_names) else f"Missing sheets: {sorted(REQUIRED_SHEETS - sheet_names)}"},
        {"id": "formula-reference", "passed": calculation_error is None, "message": calculation_error or "Required five-year DCF formulas and numeric inputs are valid"},
        {"id": "valuation-calculation", "passed": calculation_error is None, "message": calculation_error or f"Independently calculated enterprise values: {calculated}"},
        {"id": "three-scenarios", "passed": len(scenarios) == 3 and all(scenarios) and len(set(scenarios)) == 3, "message": "Three unique named scenarios are required"},
        {"id": "source-notes", "passed": bool(values.get("Sources", {}).get("A2")) and bool(values.get("Sources", {}).get("B2")), "message": "Source labels and notes are required"},
    ]
    return {"passed": all(check["passed"] for check in checks), "checks": checks, "inspection": inspection, "calculatedValues": calculated}


def main() -> None:
    raw = sys.stdin.buffer.read(MAX_INPUT_BYTES + 1)
    if len(raw) > MAX_INPUT_BYTES:
        fail("Request exceeds 512 KiB")
        return
    try:
        request = json.loads(raw)
        if not isinstance(request, dict) or not isinstance(request.get("workspace"), str):
            raise ValueError("Invalid request")
        operation = request.get("operation")
        if operation == "create_valuation":
            result = create_valuation(request["workspace"], request.get("spec"))
        elif operation == "inspect":
            result = inspect_workbook(request["workspace"], request.get("path"))
        elif operation == "verify_valuation":
            result = verify_valuation(request["workspace"], request.get("path"))
        else:
            raise ValueError("Unknown spreadsheet operation")
        print(json.dumps({"ok": True, "result": result}, separators=(",", ":")))
    except Exception as error:  # Keep worker failures structured and bounded.
        fail(str(error))


if __name__ == "__main__":
    main()
