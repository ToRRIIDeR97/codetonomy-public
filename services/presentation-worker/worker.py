#!/usr/bin/env python3
"""Schema-driven presentation creation and verification worker."""

from __future__ import annotations

import json
import os
import shutil
import subprocess
import sys
import tempfile
from pathlib import Path
from typing import Any

from PIL import Image, ImageStat
from pptx import Presentation
from pptx.chart.data import ChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt

MAX_INPUT_BYTES = 512 * 1024
MAX_DECK_BYTES = 30 * 1024 * 1024
REQUIRED_ROLES = {"title", "summary", "sources"}


def fail(message: str) -> None:
    print(json.dumps({"ok": False, "error": message}))


def executable(env_name: str, names: tuple[str, ...], candidates: tuple[str, ...]) -> str | None:
    override = os.environ.get(env_name)
    if override and Path(override).is_file():
        return override
    for name in names:
        found = shutil.which(name)
        if found:
            return found
    for candidate in candidates:
        if Path(candidate).is_file():
            return candidate
    chocolatey = os.environ.get("ChocolateyInstall")
    if chocolatey and env_name == "CODETONOMY_PDFTOPPM_BIN":
        matches = sorted(Path(chocolatey, "lib").glob("poppler*/tools/**/pdftoppm.exe"))
        if matches:
            return str(matches[0])
    local_app_data = os.environ.get("LOCALAPPDATA")
    if local_app_data and env_name == "CODETONOMY_PDFTOPPM_BIN":
        matches = sorted(Path(local_app_data, "Microsoft", "WinGet", "Packages").glob("oschwartz10612.Poppler*/**/pdftoppm.exe"))
        if matches:
            return str(matches[0])
    return None


def presentation_prerequisites() -> tuple[str, str]:
    home = Path(os.environ.get("CODETONOMY_HOME", Path.cwd() / ".codetonomy"))
    soffice = executable(
        "CODETONOMY_SOFFICE_BIN",
        ("soffice", "libreoffice"),
        (
            str(home / "prerequisites" / "libreoffice" / "program" / "soffice.exe"),
            r"C:\Program Files\LibreOffice\program\soffice.exe",
            r"C:\Program Files (x86)\LibreOffice\program\soffice.exe",
            "/Applications/LibreOffice.app/Contents/MacOS/soffice",
        ),
    )
    pdftoppm = executable(
        "CODETONOMY_PDFTOPPM_BIN",
        ("pdftoppm",),
        (
            str(home / "prerequisites" / "poppler" / "Library" / "bin" / "pdftoppm.exe"),
            str(home / "prerequisites" / "poppler" / "bin" / "pdftoppm"),
            r"C:\ProgramData\chocolatey\bin\pdftoppm.exe",
            "/opt/homebrew/bin/pdftoppm",
            "/usr/local/bin/pdftoppm",
        ),
    )
    if not soffice or not pdftoppm:
        raise ValueError(
            "Rendered slide review requires LibreOffice and pdftoppm; install both "
            "or set CODETONOMY_SOFFICE_BIN and CODETONOMY_PDFTOPPM_BIN"
        )
    return soffice, pdftoppm


def text(value: Any, label: str, maximum: int) -> str:
    if not isinstance(value, str) or not value.strip() or len(value) > maximum:
        raise ValueError(f"{label} must be non-empty text of at most {maximum} characters")
    clean = value.strip()
    if any(ord(char) < 32 and char not in "\t\n" for char in clean):
        raise ValueError(f"{label} contains unsupported control characters")
    return clean


def confined(workspace: str, requested: str, must_exist: bool) -> tuple[Path, str]:
    root = Path(workspace).resolve(strict=True)
    raw = Path(requested)
    if raw.is_absolute() or ".." in raw.parts or not raw.parts:
        raise ValueError("Presentation path is outside the workspace")
    target = root.joinpath(raw)
    ancestor = target if must_exist else target.parent
    while not ancestor.exists():
        if must_exist or ancestor == root.parent:
            raise ValueError("Presentation path does not exist")
        ancestor = ancestor.parent
    if not ancestor.resolve(strict=True).is_relative_to(root):
        raise ValueError("Presentation path resolves outside the workspace")
    if must_exist:
        if target.is_symlink() or not target.is_file() or target.stat().st_nlink != 1:
            raise ValueError("Presentation must be a regular standalone workspace file")
        if target.stat().st_size > MAX_DECK_BYTES:
            raise ValueError("Presentation exceeds 30 MiB")
    return target, raw.as_posix()


def validate_spec(raw: Any) -> dict[str, Any]:
    if not isinstance(raw, dict):
        raise ValueError("spec must be an object")
    output = text(raw.get("outputPath"), "outputPath", 1024)
    if not output.lower().endswith(".pptx"):
        raise ValueError("outputPath must end in .pptx")
    slides = raw.get("slides")
    if not isinstance(slides, list) or not 3 <= len(slides) <= 40:
        raise ValueError("slides must contain 3-40 slides")
    clean_slides = []
    roles: set[str] = set()
    for index, slide in enumerate(slides):
        if not isinstance(slide, dict):
            raise ValueError(f"slides[{index}] must be an object")
        role = slide.get("role")
        if role not in {"title", "content", "summary", "sources"}:
            raise ValueError(f"slides[{index}].role is invalid")
        roles.add(role)
        bullets = slide.get("bullets", [])
        if not isinstance(bullets, list) or len(bullets) > 10:
            raise ValueError(f"slides[{index}].bullets must contain at most 10 items")
        clean_bullets = [text(item, f"slides[{index}].bullets", 300) for item in bullets]
        chart = slide.get("chart")
        clean_chart = None
        if chart is not None:
            if not isinstance(chart, dict):
                raise ValueError(f"slides[{index}].chart must be an object")
            categories = chart.get("categories")
            series = chart.get("series")
            if not isinstance(categories, list) or not 1 <= len(categories) <= 20 or not isinstance(series, list) or not 1 <= len(series) <= 6:
                raise ValueError("chart categories and series are outside supported bounds")
            clean_categories = [text(item, "chart category", 80) for item in categories]
            clean_series = []
            for item in series:
                if not isinstance(item, dict) or not isinstance(item.get("values"), list) or len(item["values"]) != len(clean_categories):
                    raise ValueError("chart series values must match categories")
                values = item["values"]
                if any(isinstance(value, bool) or not isinstance(value, (int, float)) for value in values):
                    raise ValueError("chart values must be numbers")
                clean_series.append({"name": text(item.get("name"), "chart series name", 80), "values": values})
            clean_chart = {
                "title": text(chart.get("title"), "chart title", 160),
                "categories": clean_categories,
                "series": clean_series,
                "source": text(chart.get("source"), "chart source", 500),
            }
        clean_slides.append({
            "role": role,
            "title": text(slide.get("title"), f"slides[{index}].title", 160),
            "bullets": clean_bullets,
            "source": text(slide["source"], f"slides[{index}].source", 500) if slide.get("source") else None,
            "chart": clean_chart,
        })
    if not REQUIRED_ROLES.issubset(roles):
        raise ValueError("slides must include title, summary, and sources roles")
    return {
        "title": text(raw.get("title"), "title", 160),
        "subtitle": text(raw["subtitle"], "subtitle", 300) if raw.get("subtitle") else None,
        "outputPath": output,
        "slides": clean_slides,
    }


def add_text(slide: Any, left: float, top: float, width: float, height: float, value: str, size: int, bold: bool = False, color: str = "16324F", align: Any = None, name: str | None = None) -> Any:
    shape = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    if name:
        shape.name = name
    frame = shape.text_frame
    frame.clear()
    frame.word_wrap = True
    paragraph = frame.paragraphs[0]
    paragraph.text = value
    paragraph.font.size = Pt(size)
    paragraph.font.bold = bold
    paragraph.font.color.rgb = RGBColor.from_string(color)
    if align is not None:
        paragraph.alignment = align
    return shape


def add_bullets(slide: Any, bullets: list[str], left: float, width: float) -> None:
    shape = slide.shapes.add_textbox(Inches(left), Inches(1.55), Inches(width), Inches(5.1))
    shape.name = "codetonomy-bullets"
    frame = shape.text_frame
    frame.clear()
    frame.word_wrap = True
    for index, value in enumerate(bullets):
        paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        paragraph.text = value
        paragraph.level = 0
        paragraph.font.size = Pt(20)
        paragraph.font.color.rgb = RGBColor(0x23, 0x3D, 0x56)
        paragraph.space_after = Pt(10)


def create_deck(workspace: str, raw_spec: Any) -> dict[str, Any]:
    spec = validate_spec(raw_spec)
    target, relative_path = confined(workspace, spec["outputPath"], False)
    target.parent.mkdir(parents=True, exist_ok=True)
    deck = Presentation()
    deck.slide_width = Inches(13.333)
    deck.slide_height = Inches(7.5)
    deck.core_properties.title = spec["title"]
    deck.core_properties.subject = "Codetonomy presentation schema v1"
    while deck.slides:
        relation_id = deck.slides._sldIdLst[0].rId  # type: ignore[attr-defined]
        deck.part.drop_rel(relation_id)
        del deck.slides._sldIdLst[0]  # type: ignore[attr-defined]

    for slide_spec in spec["slides"]:
        slide = deck.slides.add_slide(deck.slide_layouts[6])
        role = slide_spec["role"]
        if role == "title":
            add_text(slide, 0.8, 1.7, 11.75, 1.4, slide_spec["title"], 34, True, align=PP_ALIGN.CENTER, name="codetonomy-role-title")
            if spec["subtitle"]:
                add_text(slide, 1.5, 3.3, 10.3, 0.9, spec["subtitle"], 20, color="4F6B82", align=PP_ALIGN.CENTER)
        else:
            add_text(slide, 0.65, 0.35, 12.0, 0.8, slide_spec["title"], 28, True, name=f"codetonomy-role-{role}")
            if slide_spec["chart"]:
                add_bullets(slide, slide_spec["bullets"], 0.7, 4.35)
                chart_spec = slide_spec["chart"]
                chart_data = ChartData()
                chart_data.categories = chart_spec["categories"]
                for series in chart_spec["series"]:
                    chart_data.add_series(series["name"], series["values"])
                chart = slide.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(5.35), Inches(1.5), Inches(7.25), Inches(4.95), chart_data).chart
                chart.has_title = True
                chart.chart_title.text_frame.text = chart_spec["title"]
                add_text(slide, 5.35, 6.62, 7.2, 0.24, f"Chart source: {chart_spec['source']}", 9, color="667788", name="codetonomy-chart-source")
            else:
                add_bullets(slide, slide_spec["bullets"], 0.9, 11.5)
        if slide_spec["source"]:
            add_text(slide, 0.7, 7.02, 11.9, 0.2, f"Source: {slide_spec['source']}", 8, color="667788", name="codetonomy-source")

    with tempfile.NamedTemporaryFile(dir=target.parent, prefix=f".{target.name}.", suffix=".tmp", delete=False) as handle:
        temporary = Path(handle.name)
    try:
        deck.save(temporary)
        with temporary.open("r+b") as handle:
            os.fsync(handle.fileno())
        os.replace(temporary, target)
    finally:
        temporary.unlink(missing_ok=True)
    return inspect_deck(workspace, relative_path, render=False)


def overlap(a: Any, b: Any) -> bool:
    horizontal = min(a.left + a.width, b.left + b.width) - max(a.left, b.left)
    vertical = min(a.top + a.height, b.top + b.height) - max(a.top, b.top)
    return horizontal > Inches(0.05) and vertical > Inches(0.05)


def likely_overflow(shape: Any) -> bool:
    sizes = [paragraph.font.size.pt for paragraph in shape.text_frame.paragraphs if paragraph.font.size]
    font_size = min(sizes, default=18)
    width = max(shape.width / Inches(1), 0.1)
    height = max(shape.height / Inches(1), 0.1)
    characters_per_line = width * 72 / (font_size * 0.55)
    lines = max(1, height * 72 / (font_size * 1.2))
    return len(shape.text) > characters_per_line * lines * 0.9


def render_deck(path: Path) -> int:
    soffice, pdftoppm = presentation_prerequisites()
    with tempfile.TemporaryDirectory(prefix="codetonomy-slides-") as directory:
        env = {
            "PATH": os.environ.get("PATH", ""),
            "HOME": str(Path.home()),
            "TMPDIR": directory,
            "XDG_CACHE_HOME": directory,
            "LANG": os.environ.get("LANG", "C.UTF-8"),
            "LC_ALL": os.environ.get("LC_ALL", "C.UTF-8"),
            "SAL_USE_VCLPLUGIN": "svp",
        }
        for name in ("SYSTEMROOT", "WINDIR", "TEMP", "TMP", "USERPROFILE", "APPDATA", "LOCALAPPDATA"):
            if name in os.environ:
                env[name] = os.environ[name]
        profile = Path(directory, "profile").as_uri()
        conversion = subprocess.run([soffice, f"-env:UserInstallation={profile}", "--headless", "--convert-to", "pdf", "--outdir", directory, str(path)], capture_output=True, timeout=45, env=env, check=False)
        pdf = Path(directory, f"{path.stem}.pdf")
        if conversion.returncode != 0 or not pdf.is_file():
            raise ValueError("LibreOffice could not render the presentation")
        raster = subprocess.run([pdftoppm, "-png", "-r", "72", str(pdf), str(Path(directory, "slide"))], capture_output=True, timeout=45, env=env, check=False)
        if raster.returncode != 0:
            raise ValueError("Rendered slides could not be rasterized")
        images = sorted(Path(directory).glob("slide-*.png"))
        for image_path in images:
            with Image.open(image_path) as image:
                grayscale = image.convert("L")
                stats = ImageStat.Stat(grayscale)
                if not stats.var or stats.var[0] < 2:
                    raise ValueError(f"Rendered slide {image_path.name} appears blank")
        return len(images)


def inspect_deck(workspace: str, requested: str, render: bool = True) -> dict[str, Any]:
    target, relative_path = confined(workspace, requested, True)
    deck = Presentation(target)
    slides = []
    collisions = 0
    overflow = 0
    roles: set[str] = set()
    charts = 0
    sourced_charts = 0
    for number, slide in enumerate(deck.slides, 1):
        title = ""
        slide_roles = []
        text_shapes = []
        for shape in slide.shapes:
            if shape.name.startswith("codetonomy-role-"):
                role = shape.name.removeprefix("codetonomy-role-")
                roles.add(role)
                slide_roles.append(role)
                if getattr(shape, "has_text_frame", False):
                    title = shape.text.strip()
            if getattr(shape, "has_chart", False):
                charts += 1
            if shape.name == "codetonomy-chart-source" and getattr(shape, "has_text_frame", False) and shape.text.strip():
                sourced_charts += 1
            if getattr(shape, "has_text_frame", False) and shape.text.strip():
                text_shapes.append(shape)
                if likely_overflow(shape):
                    overflow += 1
        for left in range(len(slide.shapes)):
            for right in range(left + 1, len(slide.shapes)):
                if overlap(slide.shapes[left], slide.shapes[right]):
                    collisions += 1
        slides.append({"number": number, "title": title, "roles": slide_roles, "shapes": len(slide.shapes), "textCharacters": sum(len(shape.text) for shape in text_shapes)})
    rendered = render_deck(target) if render else 0
    return {
        "path": relative_path,
        "slides": slides,
        "roles": sorted(roles),
        "overflowCount": overflow,
        "collisionCount": collisions,
        "chartCount": charts,
        "sourcedChartCount": sourced_charts,
        "renderedSlides": rendered,
    }


def verify_deck(workspace: str, requested: str) -> dict[str, Any]:
    inspection = inspect_deck(workspace, requested, render=True)
    roles = set(inspection["roles"])
    slide_count = len(inspection["slides"])
    checks = [
        {"id": "presentation-open", "passed": slide_count > 0, "message": f"Presentation opens with {slide_count} slides"},
        {"id": "required-slide-check", "passed": REQUIRED_ROLES.issubset(roles), "message": "Required title, summary, and sources slides are present" if REQUIRED_ROLES.issubset(roles) else "Required slide roles are missing"},
        {"id": "overflow-detection", "passed": inspection["overflowCount"] == 0, "message": f"Detected {inspection['overflowCount']} likely text overflows"},
        {"id": "layout-collision-check", "passed": inspection["collisionCount"] == 0, "message": f"Detected {inspection['collisionCount']} layout collisions"},
        {"id": "chart-source-check", "passed": inspection["chartCount"] == inspection["sourcedChartCount"], "message": f"{inspection['sourcedChartCount']} of {inspection['chartCount']} charts include sources"},
        {"id": "rendered-slide-review", "passed": inspection["renderedSlides"] == slide_count, "message": f"Rendered {inspection['renderedSlides']} of {slide_count} slides"},
    ]
    return {"passed": all(item["passed"] for item in checks), "checks": checks, "inspection": inspection}


def main() -> None:
    if sys.argv[1:] == ["--self-test"]:
        try:
            soffice, pdftoppm = presentation_prerequisites()
            print(json.dumps({"ok": True, "soffice": soffice, "pdftoppm": pdftoppm}, separators=(",", ":")))
        except Exception as error:
            print(str(error), file=sys.stderr)
            raise SystemExit(1) from error
        return
    raw = sys.stdin.buffer.read(MAX_INPUT_BYTES + 1)
    if len(raw) > MAX_INPUT_BYTES:
        fail("Request exceeds 512 KiB")
        return
    try:
        request = json.loads(raw)
        if not isinstance(request, dict) or not isinstance(request.get("workspace"), str):
            raise ValueError("Invalid request")
        operation = request.get("operation")
        if operation == "create":
            result = create_deck(request["workspace"], request.get("spec"))
        elif operation == "inspect":
            result = inspect_deck(request["workspace"], request.get("path"), render=False)
        elif operation == "verify":
            result = verify_deck(request["workspace"], request.get("path"))
        else:
            raise ValueError("Unknown presentation operation")
        print(json.dumps({"ok": True, "result": result}, separators=(",", ":")))
    except Exception as error:
        fail(str(error))


if __name__ == "__main__":
    main()
