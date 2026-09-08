"""Worker regressions; invoked by audit-regressions.test.ts with the artifact Python."""
import importlib.util
import tempfile
from pathlib import Path
from zipfile import ZipFile, ZIP_DEFLATED

from openpyxl import Workbook, load_workbook


def module(name, service):
    spec = importlib.util.spec_from_file_location(name, Path(__file__).resolve().parents[1] / 'services' / service / 'worker.py')
    loaded = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(loaded)
    return loaded


spreadsheet = module('spreadsheet', 'spreadsheet-worker')
backtest = module('backtest', 'backtesting-worker')
with tempfile.TemporaryDirectory() as directory:
    root = Path(directory)
    sparse = Workbook()
    sparse.active['A1'] = 'first'
    sparse.active['GR200'] = 'last'
    sparse.save(root / 'sparse.xlsx')
    assert spreadsheet.inspect_workbook(directory, 'sparse.xlsx')['sheets'][0]['rows'] == 200
    sparse.active['XFD1048576'] = 'far away'
    sparse.save(root / 'oversized.xlsx')
    for operation in (spreadsheet.inspect_workbook, spreadsheet.verify_valuation):
        try:
            operation(directory, 'oversized.xlsx')
            raise AssertionError('Unbounded sparse range was accepted')
        except ValueError as error:
            assert 'inspected cells' in str(error), str(error)
    # A modest compressed archive whose expansion exceeds the fixed input budget.
    with ZipFile(root / 'expanded.xlsx', 'w', ZIP_DEFLATED) as archive:
        archive.writestr('large.xml', b'0' * (spreadsheet.MAX_EXPANDED_BYTES + 1))
    try:
        spreadsheet.inspect_workbook(directory, 'expanded.xlsx')
        raise AssertionError('Expanded archive was accepted')
    except ValueError as error:
        assert 'inspection budget' in str(error)

    spreadsheet.create_valuation(directory, {
        'company': 'Fixture', 'outputPath': 'model.xlsx',
        'historical': [{'year': 2025, 'revenue': 100, 'ebitda': 20, 'freeCashFlow': 20}],
        'scenarios': [{'name': name, 'revenueGrowth': 0, 'ebitdaMargin': .2} for name in ('Bear', 'Base', 'Bull')],
        'discountRate': .1, 'terminalGrowthRate': 0, 'taxRate': 0,
        'sources': [{'label': 'Fixture', 'source': 'Synthetic data'}],
    })
    checked = spreadsheet.verify_valuation(directory, 'model.xlsx')
    assert checked['passed'], checked
    # A constant $20 annual cash flow discounted at 10% has a $200 perpetuity value.
    assert all(abs(value - 200) < 1e-9 for value in checked['calculatedValues'].values()), checked
    workbook = load_workbook(root / 'model.xlsx')
    original = workbook['Valuation']['I5'].value
    for value in ('=1/0', '=Missing!A1', '=SUM(C5:G5)', 200):
        workbook['Valuation']['I5'] = value
        workbook.save(root / 'invalid.xlsx')
        assert not spreadsheet.verify_valuation(directory, 'invalid.xlsx')['passed'], value
    workbook['Valuation']['I5'] = original
    workbook['Assumptions']['B4'] = 0
    workbook.save(root / 'invalid.xlsx')
    assert not spreadsheet.verify_valuation(directory, 'invalid.xlsx')['passed']
    workbook.close()

    spec = backtest.validate_spec({'dataPath': 'prices.csv', 'outputPath': 'backtest.json', 'shortWindow': 5,
                                  'longWindow': 30, 'trainFraction': .7, 'commissionBps': 0, 'initialCapital': 1000})
    rows = [{'date': f'2025-{index:03}', 'open': 100, 'close': 100, 'symbol': 'FIXTURE'} for index in range(33)]
    for count in (20, 31, 32):
        try:
            backtest.calculate(rows[:count], spec, 'fixture')
            raise AssertionError('Insufficient history was accepted')
        except ValueError as error:
            assert 'longWindow + 3' in str(error), str(error)
    assert backtest.calculate(rows, spec, 'fixture')['split']['index'] == 31
print('Sparse bounds, DCF formulas and calculation, and backtest window checks passed')
